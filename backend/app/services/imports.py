"""Bounded, non-executing CSV/XLSX import and preview validation."""
import csv
import re
from collections import Counter
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import PurePath
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from fastapi import HTTPException
from pydantic import ValidationError
from app.schemas import PersonalIn, CompetenciaIn
from app.services.import_normalization import HEADER_FIELDS, match_key, normalize_row

MAX_BYTES=2*1024*1024
MAX_ROWS=200


def validate_sheet_xml(archive, worksheet_path):
    """Do not trust XLSX dimension hints: reject cells outside import bounds."""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with archive.open(worksheet_path) as stream:
        for _, element in iterparse(stream, events=('end',)):
            if element.tag == ns + 'row':
                number = element.get('r', '')
                if not number.isdigit() or not 1 <= int(number) <= MAX_ROWS + 1:
                    raise ValueError('Máximo 200 filas por archivo.')
            elif element.tag == ns + 'c':
                coordinate = re.fullmatch(r'([A-T])([1-9][0-9]*)', element.get('r', ''))
                if not coordinate or int(coordinate[2]) > MAX_ROWS + 1:
                    raise ValueError('Máximo 200 filas y 20 columnas por archivo.')
            element.clear()

HEADERS={
 'personal':['cuip','curp','nombre_completo','sexo','corporacion','adscripcion','cargo','grado','estatus'],
 'competencias_basicas':['cuip','curp','institucion_evaluadora','fecha_certificacion','resultado','folio']
}
REQUIRED={'personal':{'nombre_completo','corporacion','adscripcion'},
          'competencias_basicas':{'institucion_evaluadora','fecha_certificacion','resultado','folio'}}

def target_valid(target):
    if target not in HEADERS: raise HTTPException(422,'Destino no admitido.')

def cell_text(value):
    if value is None: return ''
    if isinstance(value,(date,datetime)): return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value).strip()

def csv_dialect(text):
    try: return csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
    except csv.Error:
        # Reports often start with titles of a different width. Sniffer requires
        # consistent widths; score parsed candidate headers and repeated widths.
        def score(delimiter):
            sample=list(csv.reader(StringIO(text[:4096]),delimiter=delimiter))
            known=max((sum(match_key(cell_text(v)) in HEADER_FIELDS for v in row) for row in sample),default=0)
            widths=Counter(len(row) for row in sample if len(row)>1)
            frequency,width=max(((count,width) for width,count in widths.items()),default=(0,0))
            return known,frequency,width
        selected=max(',;\t',key=score)
        class Dialect(csv.excel): delimiter=selected
        return Dialect

def read_table(content: bytes, filename: str, sheet_name=None):
    """Read bounded raw cells without evaluating formulas or using cached values."""
    if not content or len(content)>MAX_BYTES: raise HTTPException(413,'El archivo debe contener datos y pesar como máximo 2 MB.')
    suffix=PurePath(filename.lower()).suffix
    try:
        sheets=[]; selected=None
        if suffix=='.csv':
            try: text=content.decode('utf-8-sig')
            except UnicodeDecodeError: text=content.decode('cp1252')
            if '\x00' in text: raise ValueError('Archivo CSV inválido.')
            dialect=csv_dialect(text)
            grid=[]; numbers=[]
            reader=csv.reader(StringIO(text),dialect)
            for row in reader:
                if len(row)>20: raise ValueError('Máximo 20 columnas por archivo.')
                if any(cell_text(x) for x in row):
                    grid.append(row); numbers.append(reader.line_num)
                if len(grid)>MAX_ROWS+1: raise ValueError('Máximo 200 filas por archivo. Divida la carga en lotes.')
        elif suffix=='.xlsx':
            with ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>20*1024*1024 or len(archive.infolist())>2000:
                    raise ValueError('El libro excede el tamaño permitido al descomprimir.')
                if any('vbaproject' in i.filename.lower() for i in archive.infolist()): raise ValueError('No se admiten macros.')
            from openpyxl import load_workbook
            book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                sheets=[s.title for s in book.worksheets]
                if sheet_name and sheet_name not in sheets: raise ValueError('La hoja seleccionada no existe en el libro.')
                if not sheet_name and len(sheets)>1:
                    # Inspection lists sheets first, including cover/empty sheets;
                    # importing a multi-sheet book always requires a selection.
                    return [],[],sheets,None
                sheet=book[sheet_name] if sheet_name else book.worksheets[0]
                selected=sheet.title
                if (sheet.max_row or 0)>MAX_ROWS+1 or (sheet.max_column or 0)>20: raise ValueError('Máximo 200 filas y 20 columnas por archivo.')
                with ZipFile(BytesIO(content)) as archive:
                    # Use the part resolved by openpyxl, including nonstandard paths.
                    validate_sheet_xml(archive, sheet._worksheet_path)
                # Some spreadsheet writers report A1:A1 despite having more cells.
                sheet.reset_dimensions()
                grid=[]; numbers=[]
                for n,row in enumerate(sheet.iter_rows(max_col=20),1):
                    values=[]
                    for c in row:
                        value=cell_text(c.value)
                        # Recover zeros only when an explicit integer Excel format
                        # carries them. Never invent zeros from an identifier length.
                        if c.data_type=='n' and isinstance(c.value,(int,float)) and not isinstance(c.value,bool):
                            if abs(c.value)>=10**15:
                                raise ValueError(f'Fila {n}: número con 16 o más dígitos; capture identificadores como texto para evitar pérdida de precisión.')
                            if re.fullmatch(r'0+',c.number_format) and float(c.value).is_integer() and c.value>=0:
                                value=value.zfill(len(c.number_format))
                        values.append(value)
                    if any(values): grid.append(values); numbers.append(n)
                width=max((i+1 for row in grid for i,v in enumerate(row) if v),default=0)
                grid=[row[:width] for row in grid]
            finally: book.close()
        else: raise ValueError('Use un archivo .xlsx o .csv.')
        if any(len(cell_text(v))>5000 for row in grid for v in row): raise ValueError('Una celda excede 5000 caracteres.')
        return grid,numbers,sheets,selected
    except HTTPException: raise
    except (ValueError,KeyError,csv.Error,BadZipFile,OSError,UnicodeError,ParseError,DefusedXmlException,IndexError,OverflowError) as e:
        raise HTTPException(422,str(e) if isinstance(e,ValueError) else 'El archivo está dañado o tiene un formato inválido.')


def analyze_table(content, filename, target, sheet_name=None, header_row=None):
    target_valid(target)
    grid,numbers,sheets,selected=read_table(content,filename,sheet_name)
    if not grid:
        if header_row is not None and (selected or not sheets): raise HTTPException(422,'La fila de encabezados está vacía o no existe.')
        return {'hojas':sheets,'hoja':selected,'fila_encabezado':1,'columnas':[],
                'campos':HEADERS[target],'total_filas':0,'muestra':[]},[],[]
    if header_row is not None:
        if header_row not in numbers: raise HTTPException(422,'La fila de encabezados está vacía o no existe.')
        position=numbers.index(header_row)
    else:
        def score(row):
            fields={HEADER_FIELDS.get(match_key(cell_text(v))) for v in row} & set(HEADERS[target])
            return len(fields)+ (20 if fields & {'cuip','curp'} else 0)
        position=max(range(min(10,len(grid))),key=lambda i:score(grid[i]))
    headers=[cell_text(v) for v in grid[position]]
    if len(headers)>20: raise HTTPException(422,'Máximo 20 columnas por archivo.')
    columns=[]
    for index,title in enumerate(headers):
        field=HEADER_FIELDS.get(match_key(title))
        columns.append({'indice':index,'encabezado':title or f'Columna {index+1}',
                        'campo':field if field in HEADERS[target] else None})
    analysis={'hojas':sheets,'hoja':selected,'fila_encabezado':numbers[position],
              'columnas':columns,'campos':HEADERS[target],'total_filas':len(grid)-position-1,
              'muestra':[[cell_text(v) for v in row] for row in grid[position+1:position+4]]}
    return analysis,grid[position+1:],numbers[position+1:]


def inspect_file(content, filename, target, sheet_name=None, header_row=None):
    return analyze_table(content,filename,target,sheet_name,header_row)[0]


def read_file(content: bytes, filename: str, target: str, *, sheet_name=None, header_row=None, mapping=None, with_analysis=False):
    analysis,grid,numbers=analyze_table(content,filename,target,sheet_name,header_row)
    if len(analysis['hojas'])>1 and not sheet_name: raise HTTPException(422,'Seleccione la hoja que desea importar.')
    if not grid: raise HTTPException(422,'El archivo debe incluir encabezados y al menos una fila de datos.')
    columns=analysis['columnas']
    if mapping is not None:
        if not isinstance(mapping,dict): raise HTTPException(422,'El mapeo debe relacionar campos con índices de columna.')
        if any(k not in HEADERS[target] or type(v) is not int or not 0<=v<len(columns) for k,v in mapping.items()):
            raise HTTPException(422,'El mapeo contiene un campo o índice no admitido.')
        if len(set(mapping.values()))!=len(mapping): raise HTTPException(422,'Una columna no puede alimentar dos campos.')
    else:
        fields=[c['campo'] for c in columns if c['campo']]
        if len(set(fields))!=len(fields): raise HTTPException(422,'Varias columnas corresponden al mismo campo. Ajuste el mapeo.')
        mapping={c['campo']:c['indice'] for c in columns if c['campo']}
    if not {'cuip','curp'} & set(mapping): raise HTTPException(422,'Incluya y relacione una columna CUIP o CURP.')
    if target!='personal':
        missing=REQUIRED[target]-set(mapping)
        if missing: raise HTTPException(422,'Faltan columnas: '+', '.join(sorted(missing)))
    ignored=[c['encabezado'] for c in columns if c['indice'] not in mapping.values()]
    rows=[]
    for number,row in zip(numbers,grid):
        if len(row)>len(columns) and any(cell_text(v) for v in row[len(columns):]):
            raise HTTPException(422,f'Fila {number}: hay datos fuera de las columnas de encabezados.')
        values={field:cell_text(row[index]) if index<len(row) else '' for field,index in mapping.items()}
        if not any(values.values()): raise HTTPException(422,f'Fila {number}: no contiene datos en las columnas relacionadas.')
        if any(cell_text(v).startswith(('=','+','@')) for i,v in enumerate(row) if i not in mapping.values()):
            raise HTTPException(422,f'Fila {number}: una columna omitida contiene una fórmula o prefijo no permitido.')
        rows.append((number,values))
    analysis.update(mapeo=mapping,columnas_omitidas=ignored)
    analysis.pop('muestra',None)  # Do not persist cells from ignored columns.
    return (rows,analysis) if with_analysis else rows


def build_preview(rows, target, repo, mode='combinar', analysis=None):
    if mode not in ('combinar','anexar','actualizar'): raise HTTPException(422,'Modo de carga no admitido.')
    if target!='personal' and mode!='combinar': raise HTTPException(422,'Las competencias se anexan conservando el historial. Seleccione el modo combinado.')
    original_rows=rows
    rows=[(number,normalize_row(raw)) for number,raw in rows]
    originals=dict(original_rows)
    ids={k:{r.get(k,'').upper() for _,r in rows if re.fullmatch(r'[A-Za-z0-9]{1,20}',r.get(k,''))} for k in ('cuip','curp')}
    people=repo.lookup_people(ids['cuip'],ids['curp'])
    indexes={k:{p[k]:p for p in people if p.get(k)} for k in ('cuip','curp')}
    catalogs=repo.catalogs() if target=='personal' else {}
    evaluations=repo.lookup_evaluations({p['id'] for p in people}) if target!='personal' else []
    evaluation_keys={(e['personal_id'],e['fecha_certificacion'],e['folio']) for e in evaluations}
    duplicate_keys={k:Counter(r.get(k,'').upper() for _,r in rows if r.get(k)) for k in ('cuip','curp')}
    preview=[]
    for number,raw in rows:
        errors=[]; data={}; matches={}; warnings=[]; changes=[]
        for k in ('cuip','curp'):
            ident=raw.get(k,'').upper()
            if ident in indexes[k]: matches[indexes[k][ident]['id']]=indexes[k][ident]
            if ident and duplicate_keys[k][ident]>1 and target=='personal': errors.append(f'{k.upper()} duplicado dentro del archivo.')
        person=next(iter(matches.values()),None)
        if len(matches)>1: errors.append('CUIP y CURP corresponden a personas distintas.')
        if target=='personal' and mode=='anexar' and person: errors.append('El expediente ya existe. El modo anexar no permite actualizarlo.')
        if target=='personal' and mode=='actualizar' and not person: errors.append('No se encontró el expediente. El modo actualizar no permite crear registros.')
        if any(v.startswith(('=','+','@')) for v in raw.values()): errors.append('No se admiten fórmulas ni celdas que comiencen con =, + o @.')
        try:
            missing=[k for k in REQUIRED[target] if not raw.get(k)] if not person or target!='personal' else []
            if missing: raise ValueError('Campos obligatorios vacíos: '+', '.join(sorted(missing)))
            # Validate both supplied identifiers, including when importing evaluations.
            if not raw.get('cuip') and not raw.get('curp'): raise ValueError('Capture CUIP o CURP.')
            for k in ('cuip','curp'):
                if raw.get(k):
                    if k=='cuip': PersonalIn.check_cuip(raw[k].upper())
                    else: PersonalIn.check_curp(raw[k].upper())
            if target=='personal':
                # A single matching identifier must never overwrite a conflicting
                # identifier already assigned to that expediente.
                if person and any(raw.get(k) and person.get(k) and person[k]!=raw[k] for k in ('cuip','curp')):
                    raise ValueError('Un identificador proporcionado difiere del expediente. Corríjalo antes de actualizar.')
                baseline={k:person[k] for k in PersonalIn.model_fields if person and k in person}
                baseline.update({k:v for k,v in raw.items() if k in PersonalIn.model_fields and v})
                for src,dest in [('corporacion','corporacion_id'),('cargo','cargo_id'),('grado','grado_id')]:
                    if raw.get(src):
                        table={'corporacion':'corporaciones','cargo':'cargos','grado':'grados'}[src]
                        found=[c for c in catalogs[table] if match_key(c['nombre'])==match_key(raw[src]) and c['activo']]
                        if len(found)>1: raise ValueError(f'{src.capitalize()} tiene coincidencias ambiguas en el catálogo: {raw[src]}')
                        if not found: raise ValueError(f'{src.capitalize()} no existe o está inactivo: {raw[src]}')
                        baseline[dest]=found[0]['id']
                data=PersonalIn(**baseline).model_dump(mode='json')
                for key,value in data.items():
                    old=person.get(key) if person else None
                    if value!=old:
                        field=key.removesuffix('_id')
                        table={'corporacion':'corporaciones','cargo':'cargos','grado':'grados'}.get(field)
                        def display(v):
                            return next((c['nombre'] for c in catalogs[table] if c['id']==v),v) if table else v
                        changes.append({'campo':field,'anterior':display(old),'nuevo':display(value)})
                if person and not changes: warnings.append('El expediente ya contiene estos valores; no hay cambios de contenido.')
            else:
                if not person: raise ValueError('No se encontró un expediente con ese CUIP o CURP.')
                if any(raw.get(k) and person.get(k)!=raw[k].upper() for k in ('cuip','curp')):
                    raise ValueError('Los identificadores no coinciden con el expediente.')
                data=CompetenciaIn(**{k:v for k,v in raw.items() if k not in ('cuip','curp')},personal_id=person['id']).model_dump(mode='json')
                if (person['id'],data['fecha_certificacion'],data['folio']) in evaluation_keys:
                    errors.append('La evaluación ya existe en el expediente.')
        except ValidationError as e:
            errors.extend('.'.join(str(x) for x in err['loc'])+': '+err['msg'].replace('Value error, ','') for err in e.errors())
        except ValueError as e: errors.append(str(e))
        preview.append({'fila':number,'identificador':raw.get('cuip') or raw.get('curp') or '',
            'accion':'actualizar' if person and target=='personal' else 'insertar','personal_id':person['id'] if person else None,
            'version':person.get('updated_at') if person else None,'datos':data,'origen':originals[number],
            'normalizados':raw,'cambios':changes,'avisos':warnings,'modo':mode,'errores':errors})
    # Different identifiers can still refer to the same person.
    keys=Counter((r['personal_id'],r['datos'].get('fecha_certificacion'),r['datos'].get('folio')) if target!='personal' else r['personal_id'] for r in preview if r['personal_id'])
    for r in preview:
        key=(r['personal_id'],r['datos'].get('fecha_certificacion'),r['datos'].get('folio')) if target!='personal' else r['personal_id']
        if r['personal_id'] and keys[key]>1: r['errores'].append('Registro repetido en el archivo.')
    if analysis and preview: preview[0]['analisis']=analysis
    return preview

def template(target):
    target_valid(target)
    out=StringIO(); csv.writer(out).writerow(HEADERS[target])
    return out.getvalue().encode('utf-8-sig')
