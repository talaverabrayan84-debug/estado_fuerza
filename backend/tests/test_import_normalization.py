"""Exercise ingestion through the API and verify persisted records, not just strings."""
import csv
import json
from copy import deepcopy
from io import BytesIO, StringIO

import pytest
from openpyxl import Workbook

from app.db import demo


def csv_file(rows, delimiter=',', encoding='utf-8-sig'):
    out=StringIO(); csv.writer(out,delimiter=delimiter).writerows(rows)
    return out.getvalue().encode(encoding)


def post(client, headers, rows=None, *, content=None, filename='datos.csv', endpoint='previsualizar', **options):
    return client.post('/api/carga-masiva/'+endpoint,headers=headers,
                       data={'destino':'personal',**options},
                       files={'archivo':(filename,content if content is not None else csv_file(rows))})


def confirm(client, headers, record):
    return client.post(f"/api/carga-masiva/{record['id']}/confirmar",headers=headers)


@pytest.mark.parametrize('delimiter,encoding',[(',', 'utf-8-sig'),(';','cp1252'),('\t','utf-8')])
def test_aliases_normalization_and_explicit_ignored_columns(client,auth,delimiter,encoding):
    headers=auth()
    rows=[['Reporte de personal'],[],['CUIP','Nombre del elemento','Corporación','Unidad de adscripción','Género','Observaciones'],
          [' nuevo01 ','  Persona   importada  ','corporacion de ejemplo a','  Unidad   nueva ','femenino','No importar']]
    content=csv_file(rows,delimiter,encoding)
    before=deepcopy(demo.store.people)
    analysis=post(client,headers,content=content,endpoint='analizar')
    assert analysis.status_code==200,analysis.text
    assert analysis.json()['fila_encabezado']==3
    assert demo.store.people==before and not demo.store.imports
    response=post(client,headers,content=content)
    assert response.status_code==201,response.text
    record=response.json(); assert record['registros_error']==0
    row=record['filas'][0]
    assert row['fila']==4 and row['datos']['cuip']=='NUEVO01'
    assert row['datos']['nombre_completo']=='Persona importada' and row['datos']['sexo']=='M'
    assert row['analisis']['columnas_omitidas']==['Observaciones']
    assert 'muestra' not in row['analisis']
    assert confirm(client,headers,record).status_code==200
    person=next(p for p in demo.store.people.values() if p['cuip']=='NUEVO01')
    assert person['corporacion_id']==demo.CORP_A and person['adscripcion']=='Unidad nueva'


def test_partial_update_preserves_blank_and_absent_fields_and_reports_diff(client,auth):
    headers=auth(); before=deepcopy(demo.store.people[demo.WORKER_ID])
    record=post(client,headers,[['CUIP','Adscripción','Nombre completo'],['demo00001',' Unidad nueva ','']],modo='actualizar').json()
    assert record['registros_error']==0
    row=record['filas'][0]
    assert row['accion']=='actualizar'
    assert row['datos']['nombre_completo']==before['nombre_completo']
    assert row['cambios']==[{'campo':'adscripcion','anterior':before['adscripcion'],'nuevo':'Unidad nueva'}]
    assert confirm(client,headers,record).status_code==200
    after=demo.store.people[demo.WORKER_ID]
    assert after['adscripcion']=='Unidad nueva'
    for key in ('cuip','curp','sexo','nombre_completo','corporacion_id','cargo_id','grado_id','estatus'):
        assert after[key]==before[key]
    report=client.get(f"/api/carga-masiva/{record['id']}/reporte",headers=headers)
    assert 'cambios' in report.text and 'Unidad nueva' in report.text


@pytest.mark.parametrize('mode,identifier',[('anexar','DEMO00001'),('actualizar','NUEVO01')])
def test_modes_reject_unintended_actions_and_cannot_confirm(client,auth,mode,identifier):
    headers=auth(); before=deepcopy(demo.store.people)
    record=post(client,headers,[['cuip','nombre_completo','corporacion','adscripcion'],
                               [identifier,'Persona nueva','Corporación de ejemplo A','Unidad']],modo=mode).json()
    assert record['registros_error']==1
    assert confirm(client,headers,record).status_code==422
    assert demo.store.people==before


def test_combined_batch_updates_and_appends_idempotently(client,auth):
    headers=auth(); before=len(demo.store.people)
    record=post(client,headers,[['cuip','nombre_completo','corporacion','adscripcion'],
                               ['DEMO00001','','','Unidad cambiada'],
                               ['NUEVO01','Persona nueva','Corporación de ejemplo A','Unidad nueva']]).json()
    assert record['registros_error']==0
    assert [r['accion'] for r in record['filas']]==['actualizar','insertar']
    assert confirm(client,headers,record).status_code==200
    assert confirm(client,headers,record).status_code==200
    assert len(demo.store.people)==before+1
    assert demo.store.people[demo.WORKER_ID]['adscripcion']=='Unidad cambiada'


def test_missing_required_fields_are_still_rejected_for_new_person(client,auth):
    headers=auth()
    record=post(client,headers,[['cuip','adscripcion'],['NUEVO01','Unidad nueva']]).json()
    assert record['registros_error']==1
    assert confirm(client,headers,record).status_code==422


def test_conflicting_unmatched_identifier_cannot_overwrite_person(client,auth):
    headers=auth(); demo.store.people[demo.WORKER_ID]['curp']='AAAA900101HDFBBB01'
    before=deepcopy(demo.store.people[demo.WORKER_ID])
    record=post(client,headers,[['cuip','curp','adscripcion'],['OTROCUIP',before['curp'],'Unidad nueva']]).json()
    assert record['registros_error']==1
    assert any('identificador' in error for error in record['filas'][0]['errores'])
    assert confirm(client,headers,record).status_code==422
    assert demo.store.people[demo.WORKER_ID]==before


def test_catalog_normalization_does_not_guess_ambiguous_names(client,auth):
    headers=auth()
    demo.store.catalogs['corporaciones'].append({'id':'30000000-0000-0000-0000-000000000001',
                                               'nombre':'Corporacion de ejemplo A','activo':True})
    record=post(client,headers,[['cuip','corporacion'],['DEMO00001','Corporación de ejemplo A']]).json()
    assert record['registros_error']==1
    assert any('ambiguas' in error for error in record['filas'][0]['errores'])


def test_manual_mapping_resolves_unknown_or_duplicate_headings(client,auth):
    headers=auth()
    rows=[['Código policial','Ubicación'],['DEMO00001','Unidad manual']]
    assert post(client,headers,rows).status_code==422
    record=post(client,headers,rows,mapeo=json.dumps({'cuip':0,'adscripcion':1})).json()
    assert record['registros_error']==0
    assert confirm(client,headers,record).status_code==200
    assert demo.store.people[demo.WORKER_ID]['adscripcion']=='Unidad manual'
    rows=[['CUIP','Adscripción','Unidad'],['DEMO00001','Unidad uno','Unidad dos']]
    assert post(client,headers,rows).status_code==422
    record=post(client,headers,rows,mapeo=json.dumps({'cuip':0,'adscripcion':2})).json()
    assert record['registros_error']==0
    assert record['filas'][0]['datos']['adscripcion']=='Unidad dos'


@pytest.mark.parametrize('mapping',['null','[]','bad','{"sql":0}','{"cuip":true}',
                                    '{"cuip":20}','{"cuip":-1}','{"cuip":0,"curp":0}'])
def test_invalid_mapping_never_stages_or_writes(client,auth,mapping):
    assert post(client,auth(),[['cuip','adscripcion'],['DEMO00001','Unidad']],mapeo=mapping).status_code==422
    assert not demo.store.imports


def workbook_bytes(book):
    out=BytesIO(); book.save(out); book.close(); return out.getvalue()


def test_multiple_sheets_require_selection_even_with_empty_cover(client,auth):
    headers=auth(); book=Workbook(); book.active.title='Portada'
    sheet=book.create_sheet('Personal'); sheet.append(['CUIP','Adscripción']);sheet.append(['DEMO00001','Unidad de hoja dos'])
    content=workbook_bytes(book)
    analysis=post(client,headers,content=content,filename='datos.xlsx',endpoint='analizar').json()
    assert analysis['hojas']==['Portada','Personal'] and analysis['hoja'] is None
    assert post(client,headers,content=content,filename='datos.xlsx').status_code==422
    assert post(client,headers,content=content,filename='datos.xlsx',hoja='Inexistente').status_code==422
    record=post(client,headers,content=content,filename='datos.xlsx',hoja='Personal').json()
    assert record['registros_error']==0
    assert record['filas'][0]['analisis']['hoja']=='Personal'
    assert confirm(client,headers,record).status_code==200
    assert demo.store.people[demo.WORKER_ID]['adscripcion']=='Unidad de hoja dos'


def test_excel_zero_format_preserves_identifier_and_large_numbers_reject(client,auth):
    headers=auth(); book=Workbook();sheet=book.active
    sheet.append(['cuip','nombre_completo','corporacion','adscripcion'])
    sheet.append([123,'Persona con ceros','Corporación de ejemplo A','Unidad'])
    sheet['A2'].number_format='000000'
    record=post(client,headers,content=workbook_bytes(book),filename='datos.xlsx').json()
    assert record['registros_error']==0 and record['filas'][0]['datos']['cuip']=='000123'
    book=Workbook();book.active.append(['cuip','adscripcion']);book.active.append([1234567890123456,'Unidad'])
    assert post(client,headers,content=workbook_bytes(book),filename='datos.xlsx').status_code==422


@pytest.mark.parametrize('date_value,valid',[('31/01/2020',True),('31-01-2020',True),('01/02/2020',True),
                                         ('2020-01-31',True),('31/02/2020',False),('0',False),
                                         ('43831',False),('2020-01-01T00:00:00Z',False)])
def test_evaluation_dates_and_results(client,auth,date_value,valid):
    headers=auth()
    record=post(client,headers,[['CUIP','Institución evaluadora','Fecha de certificación','Resultado','Folio'],
                               ['demo00001','Academia ficticia',date_value,'NO APROBADO','FOLIO-N']],
                destino='competencias_basicas').json()
    assert record['registros_error']==(0 if valid else 1)
    if valid:
        assert record['filas'][0]['datos']['resultado']=='no_aprobado'
        if date_value=='01/02/2020': assert record['filas'][0]['datos']['fecha_certificacion']=='2020-02-01'
        assert confirm(client,headers,record).status_code==200
        assert any(e['folio']=='FOLIO-N' for e in demo.store.evaluations)
    else: assert confirm(client,headers,record).status_code==422


def test_formula_in_ignored_column_and_unauthorized_analysis_are_rejected(client,auth):
    rows=[['cuip','adscripcion','Observaciones'],['DEMO00001','Unidad','=SUM(1,2)']]
    assert post(client,auth(),rows).status_code==422
    assert post(client,auth('trabajador'),rows,endpoint='analizar').status_code==403
    assert post(client,{},rows,endpoint='analizar').status_code==401


def test_explicit_header_row_and_trailing_empty_csv_cells(client,auth):
    rows=[['Reporte'],['Encabezado uno'],['CUIP','Adscripción','Cargo'],['DEMO00001','Unidad']]
    record=post(client,auth(),rows,fila_encabezado='3').json()
    assert record['registros_error']==0 and record['filas'][0]['fila']==4
    assert record['filas'][0]['origen']['cargo']==''
    assert post(client,auth(),rows,fila_encabezado='200').status_code==422


def test_unknown_mode_does_not_stage_batch(client,auth):
    assert post(client,auth(),[['cuip','adscripcion'],['DEMO00001','Unidad']],modo='borrar').status_code==422
    assert not demo.store.imports
