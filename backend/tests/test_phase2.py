import csv
from io import BytesIO, StringIO
from zipfile import ZipFile
import re
from datetime import timedelta
import pytest
from openpyxl import Workbook
from app.db import demo
from app.services.vigencia import hoy_local

def course(client,h):
    r=client.post('/api/cursos',headers=h,json=dict(nombre_curso='Curso de prueba',tipo='competencias_basicas',fuente_financiamiento='FASP',institucion_impartidora='Academia ficticia',horas=40))
    assert r.status_code==201,r.text
    assert r.json()['vigencia_meses']==36
    return r.json()['id']

def session(client,h,cupo=1):
    r=client.post('/api/curso-sesiones',headers=h,json=dict(curso_id=course(client,h),fecha_inicio=hoy_local().isoformat(),fecha_fin=hoy_local().isoformat(),sede='Aula prueba',modalidad='presencial',cupo=cupo))
    assert r.status_code==201,r.text
    return r.json()

def csv_bytes(rows):
    s=StringIO();csv.writer(s).writerows(rows);return s.getvalue().encode('utf-8-sig')

HEAD=['cuip','nombre_completo','corporacion','adscripcion']
ROW=['NUEVO001','Persona importada','Corporación de ejemplo A','Unidad importada']

def preview(client,h,rows=None,content=None,name='prueba.csv',target='personal'):
    return client.post('/api/carga-masiva/previsualizar',headers=h,data={'destino':target},files={'archivo':(name,content if content is not None else csv_bytes(rows or [HEAD,ROW]))})

def confirm(client,h,r):
    return client.post('/api/carga-masiva/'+r['id']+'/confirmar',headers=h)

def test_capacity_updates_and_worker_boundaries(client,auth):
    h=auth('capacitacion');s=session(client,h);sid=s['id'];p,q=list(demo.store.people)[:2]
    path=f'/api/curso-sesiones/{sid}/inscripciones'
    assert client.post(path,headers=h,json={'personal_id':p}).status_code==200
    assert client.post(path,headers=h,json={'personal_id':q}).status_code==409
    assert client.post(path,headers=h,json={'personal_id':p,'estatus':'baja'}).status_code==200
    assert client.post(path,headers=h,json={'personal_id':q}).status_code==200
    assert client.post(path,headers=h,json={'personal_id':p}).status_code==409
    s.pop('id');s['cupo']=2
    assert client.put('/api/curso-sesiones/'+sid,headers=h,json=s).status_code==200
    client.post(path,headers=h,json={'personal_id':p})
    s['cupo']=1
    assert client.put('/api/curso-sesiones/'+sid,headers=h,json=s).status_code==409
    w=auth('trabajador')
    assert client.post(path,headers=w,json={'personal_id':q}).status_code==403
    own=client.get(path,headers=w).json()['items']
    assert len(own)==1 and own[0]['personal_id']==p
    assert client.get(f'/api/personal/{q}/inscripciones',headers=w).status_code==404

def test_calendar_overlap_dates_and_pagination(client,auth):
    h=auth();s=session(client,h);d=hoy_local().isoformat()
    r=client.get(f'/api/curso-sesiones?desde={d}&hasta={d}&limit=1',headers=h)
    assert r.status_code==200 and r.json()['has_more']
    rows=client.get(f'/api/curso-sesiones?desde={d}&hasta={d}',headers=h).json()['items']
    assert any(x['id']==s['id'] and x['estatus']=='En curso' for x in rows)
    assert client.get('/api/curso-sesiones?desde=2026-01-01&hasta=2025-01-01',headers=h).status_code==422
    s.pop('id');s['fecha_fin']='2020-01-01'
    assert client.post('/api/curso-sesiones',headers=h,json=s).status_code==422

def test_journal_ownership_and_validation(client,auth):
    w=auth('trabajador');a=auth();p=demo.WORKER_ID;s=next(iter(demo.store.sessions))
    data=dict(sesion_id=s,fecha=hoy_local().isoformat(),tipo_actividad='avance',descripcion='Actividad ficticia realizada')
    assert client.post('/api/bitacora',headers=a,json=data).status_code==403
    assert client.post('/api/bitacora',headers=w,json={**data,'personal_id':p}).status_code==422
    assert client.post('/api/bitacora',headers=w,json={**data,'fecha':(hoy_local()+timedelta(days=1)).isoformat()}).status_code==422
    r=client.post('/api/bitacora',headers=w,json=data);assert r.status_code==201,r.text
    eid=r.json()['id']
    assert client.put('/api/bitacora/'+eid,headers=w,json={**data,'descripcion':'Actividad editada'}).status_code==200
    assert len(client.get(f'/api/personal/{p}/bitacora',headers=a).json()['items'])==1
    other=list(demo.store.people)[1]
    assert client.get(f'/api/personal/{other}/bitacora',headers=w).status_code==404
    demo.store.journal[eid]['personal_id']=other
    assert client.put('/api/bitacora/'+eid,headers=w,json=data).status_code==404
    unregistered=session(client,a)['id']
    assert client.post('/api/bitacora',headers=w,json={**data,'sesion_id':unregistered}).status_code==403

@pytest.mark.parametrize('ext',['csv','xlsx'])
def test_import_preview_then_idempotent_confirm(client,auth,ext):
    h=auth();before=len(demo.store.people)
    if ext=='xlsx':
        b=Workbook();b.active.append(HEAD);b.active.append(ROW);out=BytesIO();b.save(out);content=out.getvalue()
    else: content=csv_bytes([HEAD,ROW])
    r=preview(client,h,content=content,name='prueba.'+ext);assert r.status_code==201,r.text
    record=r.json();assert record['registros_error']==0 and len(demo.store.people)==before
    assert confirm(client,h,record).json()['registros_procesados']==1
    assert confirm(client,h,record).status_code==200
    assert len(demo.store.people)==before+1
    p=next(p for p in demo.store.people.values() if p['cuip']=='NUEVO001');assert p['adscripcion']=='Unidad importada'

def test_import_errors_report_and_roles(client,auth):
    h=auth();r=preview(client,h,rows=[HEAD,ROW,ROW]).json()
    assert r['registros_error']==2
    assert confirm(client,h,r).status_code==422
    assert client.get('/api/carga-masiva/'+r['id']+'/reporte',headers=h).status_code==200
    assert preview(client,auth('capacitacion')).status_code==403
    assert client.get('/api/carga-masiva',headers=auth('trabajador')).status_code==403
    assert client.get('/api/carga-masiva').status_code==401

def test_stale_and_atomic_conflicts(client,auth):
    h=auth();p=demo.store.people[demo.WORKER_ID]
    rows=[HEAD,['DEMO00001','Nombre cambiado','Corporación de ejemplo A','Unidad editada']]
    r=preview(client,h,rows=rows).json();assert r['registros_error']==0
    p['updated_at']='new version'
    assert confirm(client,h,r).status_code==409
    r=preview(client,h,rows=[HEAD,ROW,['NUEVO002',*ROW[1:]]]).json()
    # A concurrent insert after preview must roll back the earlier row too.
    demo.DemoRepository(demo.store,{'rol':'admin'}).save_personal({**{k:v for k,v in p.items() if k not in ('id','updated_at')},'cuip':'NUEVO002'})
    assert confirm(client,h,r).status_code==409
    assert not any(p['cuip']=='NUEVO001' for p in demo.store.people.values())
    assert demo.store.imports[r['id']]['estado']=='revision'

def test_invalid_files_formula_and_original_row_number(client,auth):
    h=auth()
    assert preview(client,h,content=b'bad',name='wrong.xlsx').status_code==422
    assert preview(client,h,content=b'x'*(2*1024*1024+1)).status_code==413
    assert preview(client,h,rows=[HEAD]+[ROW]*201).status_code==422
    r=preview(client,h,rows=[HEAD,[],['NUEVO001','=SUM(1)','Corporación de ejemplo A','Unidad']]).json()
    assert r['filas'][0]['fila']==3 and r['registros_error']==1
    assert preview(client,h,rows=[['bad'],['value']]).status_code==422
    r=preview(client,h,rows=[HEAD,['DEMO00001','','Corporación de ejemplo A','Unidad']]).json()
    assert r['registros_error']==1

def test_evaluations_import_duplicate_and_history(client,auth):
    h=auth();head=['cuip','institucion_evaluadora','fecha_certificacion','resultado','folio']
    row=['DEMO00001','Academia ficticia','2020-01-01','aprobado','IMPORT001']
    r=preview(client,h,rows=[head,row],target='competencias_basicas').json()
    assert r['registros_error']==0
    assert confirm(client,h,r).status_code==200
    e=next(e for e in demo.store.evaluations if e['folio']=='IMPORT001');assert not e['activo']
    r=preview(client,h,rows=[head,row],target='competencias_basicas').json();assert r['registros_error']==1

def test_templates_are_real_excel_and_csv(client,auth):
    h=auth()
    for ext in ['csv','xlsx']:
        r=client.get('/api/carga-masiva/plantilla?destino=personal&formato='+ext,headers=h)
        assert r.status_code==200
        assert r.content.startswith(b'PK' if ext=='xlsx' else b'\xef\xbb\xbf')


def test_xlsx_dimension_hints_cannot_hide_rows_or_columns(client, auth):
    h=auth()

    def workbook(cell=None):
        book=Workbook();book.active.append(HEAD);book.active.append(ROW)
        if cell: book.active[cell]='hidden'
        source=BytesIO();book.save(source);book.close()
        result=BytesIO()
        with ZipFile(BytesIO(source.getvalue())) as original, ZipFile(result,'w') as altered:
            for entry in original.infolist():
                content=original.read(entry.filename)
                if entry.filename=='xl/worksheets/sheet1.xml':
                    content=re.sub(rb'<dimension ref="[^"]+"/>',b'<dimension ref="A1:A1"/>',content)
                altered.writestr(entry,content)
        return result.getvalue()

    # A harmless incorrect hint must not discard the valid data cells.
    response=preview(client,h,content=workbook(),name='hint.xlsx')
    assert response.status_code==201,response.text
    assert response.json()['filas'][0]['datos']['cuip']=='NUEVO001'
    for cell in ['A202','U2']:
        response=preview(client,h,content=workbook(cell),name='hidden.xlsx')
        assert response.status_code==422,response.text


def test_excel_template_accepts_empty_optional_trailing_cells(client,auth):
    from openpyxl import load_workbook
    h=auth()
    template=client.get('/api/carga-masiva/plantilla?destino=personal&formato=xlsx',headers=h).content
    book=load_workbook(BytesIO(template))
    book.active.append(['NUEVO001',None,'Persona importada',None,'Corporación de ejemplo A','Unidad importada'])
    out=BytesIO();book.save(out);book.close()
    response=preview(client,h,content=out.getvalue(),name='plantilla.xlsx')
    assert response.status_code==201,response.text
    assert response.json()['registros_error']==0

def test_evidence_private_immutable_and_type_checked(client,auth):
    w=auth('trabajador');a=auth()
    r=client.post('/api/bitacora',headers=w,json={'fecha':hoy_local().isoformat(),'tipo_actividad':'entrega','descripcion':'Entrega con comprobante'}).json()
    path=f"/api/bitacora/{r['id']}/evidencia"
    assert client.post(path,headers=w,files={'archivo':('fake.pdf',b'not PDF')}).status_code==422
    assert client.post(path,headers=a,files={'archivo':('demo.pdf',b'%PDF-1.4\n%%EOF')}).status_code==403
    content=b'%PDF-1.4\n%%EOF'
    assert client.post(path,headers=w,files={'archivo':('demo.pdf',content)}).status_code==201
    assert client.post(path,headers=w,files={'archivo':('demo.pdf',content)}).status_code==409
    assert client.get(path+'/descargar').status_code==401
    assert client.get(path+'/descargar',headers=a).content==content
    assert 'attachment' in client.get(path+'/descargar',headers=w).headers['content-disposition']
    demo.store.journal[r['id']]['personal_id']=list(demo.store.people)[1]
    assert client.get(path+'/descargar',headers=w).status_code==404
